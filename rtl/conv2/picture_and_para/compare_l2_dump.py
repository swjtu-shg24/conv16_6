# -*- coding: utf-8 -*-
"""compare_l2_dump.py —— L2 **全部仿真数据**的外部对拍 + 成表

把 tb_top_l2 抓的 **真实数据通路** L2 逐级数据（`real_dump_l2.txt`）与 Python golden
（`l2_golden_real_flat.hex`，由 gen_stim.py 生成）**逐点**对拍，逐级打印统计，并写出
`feature_maps_real_l2.xlsx`（每个 tile 一个 sheet）方便一个个数字对着看。

跑法（工程根目录，conda cyclegan 环境）：
    & 'D:\\Users\\Administrator\\anaconda3\\envs\\cyclegan\\python.exe' rtl\\conv2\\picture_and_para\\compare_l2_dump.py

行序（与 tb_top_l2.v 的抓数一致，每个 tile 64 行 × 100 值）：
    0..7    DWC  (8 通道，量化后的 dw 输出)
    8..15   BNR  (8 通道，dw 侧归一化 + ReLU 之后)
    16..31  QQ   (16 oc，pw 量化)
    32..47  BNQ  (16 oc，pw 侧归一化，可负)
    48..63  POOL (16 oc，2×2 max → 5×5，不足 100 补 0)
"""
import os
import sys

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import stim_model as M

HERE = M.HERE
TILES = [(0, 0), (5, 7), (11, 15)]        # 与 tb_top_l2.v / gen_stim.py 必须一致
NL, NV = 64, 100
OUTX = os.path.join(HERE, "feature_maps_real_l2.xlsx")

# 行键：与 gen_stim.py 的行序严格对应
KEYS = ([("DWC", ch) for ch in range(8)] +
        [("BNR", ch) for ch in range(8)] +
        [("QQ", oc) for oc in range(16)] +
        [("BNQ", oc) for oc in range(16)] +
        [("POOL", oc) for oc in range(16)])
STAGES = [("DWC", 0, 8), ("BNR", 8, 16), ("QQ", 16, 32), ("BNQ", 32, 48), ("POOL", 48, 64)]


def read_rows(path):
    """读"每行若干十六进制值"→ 每行一个 list[int]，解析不了（含 x/z）的记为 None

    ★ tb 的 $fdisplay 会把源码里的中文按 UTF-8 写进去，所以按 UTF-8 读；
    ★ 含 x/z 的行**保留**（记 None），不要静默丢掉 —— 那正是"抓数没落地"的信号。
    """
    rows = []
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        for ln in f:
            ln = ln.strip()
            if (not ln) or ln.startswith("#"):
                continue
            try:
                rows.append([int(x, 16) for x in ln.split()])
            except ValueError:
                rows.append(None)
    return rows


def main():
    gp = os.path.join(HERE, "l2_golden_real_flat.hex")
    rp = os.path.join(HERE, "real_dump_l2.txt")
    for p in (gp, rp):
        if not os.path.exists(p):
            print("缺文件：%s" % p)
            print("  提示：先跑 gen_stim.py 生成 golden；再跑 vsim -c -do rtl/conv2/sim/run_l2.do 抓 RTL 数据")
            return 1

    g = read_rows(gp)
    r = read_rows(rp)
    n_tile = len(TILES)
    print("=" * 78)
    print("compare_l2_dump : L2 逐级数据 RTL vs Golden（真实图 test.jpg + 真实权重）")
    print("=" * 78)
    print("  golden 行数 %d（期望 %d）；RTL 行数 %d" % (len(g), n_tile * NL, len(r)))
    if len(r) < n_tile * NL or len(g) < n_tile * NL:
        print("  行数不足：先确认 tb_top_l2 跑完并写出了 real_dump_l2.txt")
        return 1

    total_bad = 0
    total_x = 0
    total_n = 0
    detail = []          # (ti, stage_name, n, n_x, n_bad, max|diff|)
    for ti, (tr, tc) in enumerate(TILES):
        print("  ---- tile %d (%2d,%2d) ----" % (ti, tr, tc))
        for sname, k0, k1 in STAGES:
            n = n_x = n_bad = 0
            mx = 0
            first = []
            for k in range(k0, k1):
                a = r[ti*NL + k]
                b = g[ti*NL + k]
                if a is None:
                    n_x += len(b)
                    n += len(b)
                    continue
                for i in range(min(len(a), len(b))):
                    n += 1
                    if a[i] != b[i]:
                        n_bad += 1
                        mx = max(mx, abs(a[i] - b[i]))
                        if len(first) < 4:
                            first.append((KEYS[k][1], i, a[i], b[i]))
            detail.append((ti, sname, n, n_x, n_bad, mx))
            total_n += n
            total_x += n_x
            total_bad += n_bad
            line = ("    %-4s : 点 %5d  含x %4d  不一致 %4d  最大|差| %3d" %
                    (sname, n, n_x, n_bad, mx))
            print(line)
            for (idx, i, av, bv) in first:
                print("         %s[%d] p%d: RTL %d (%.3f)  GOLD %d (%.3f)"
                      % (sname, idx, i, av, av/16.0, bv, bv/16.0))

    # ---- 成表 ----
    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws0 = wb.active
        ws0.title = "Params"
        ws0.append(["项", "值"])
        ws0.append(["数据来源", "tb_top_l2.v（真实 test.jpg + 真实 wrom.hex 315 字）"])
        ws0.append(["tile 列表", str(TILES)])
        ws0.append(["比对点数", total_n])
        ws0.append(["★ 含 x 的抓数点", total_x])
        ws0.append(["★ 不一致点数", total_bad])
        ws0.append(["判定", "PASS" if (total_bad == 0 and total_x == 0) else "FAIL"])
        ws0.append(["说明", "数值都是 Q4.4 寄存器值；实际值 = 数字 / 16"])
        ws0.append(["行序", "DWC(0..7) BNR(8..15) QQ(16..31) BNQ(32..47) POOL(48..63)"])
        ws0.append([])
        ws0.append(["tile", "级", "点数", "含x", "不一致", "最大|差|"])
        for (ti, sname, n, n_x, n_bad, mx) in detail:
            ws0.append([ti, sname, n, n_x, n_bad, mx])
        for ti, (tr, tc) in enumerate(TILES):
            ws = wb.create_sheet("L2_t%d_%d" % (tr, tc))
            ws.append(["tile %d (%d,%d)  RTL vs GOLD" % (ti, tr, tc)])
            ws.append(["级", "序号"] + ["p%d" % i for i in range(NV)])
            for k, (nm, oc) in enumerate(KEYS):
                a = r[ti*NL + k]
                b = g[ti*NL + k]
                if a is None:
                    ws.append(["%s%d" % (nm, oc), "RTL"] + ["x"] * NV)
                    ws.append(["", "GOLD"] + b)
                    ws.append(["", "DIFF"] + ["?"] * NV)
                    continue
                ws.append(["%s%d" % (nm, oc), "RTL"] + a)
                ws.append(["", "GOLD"] + b)
                ws.append(["", "DIFF"] + [int(x - y) for x, y in zip(a, b)])
        wb.save(OUTX)
        print("  写出 %s（%d 个 sheet）" % (os.path.basename(OUTX), len(wb.sheetnames)))
    except ImportError:
        print("  （没装 openpyxl，跳过 xlsx）")

    print("-" * 78)
    print("  合计：比对 %d 点，含 x %d 点，不一致 %d 点" % (total_n, total_x, total_bad))
    print("RTL vs Golden 不一致点数 = %d" % total_bad)
    # ★ ASCII 判定行（给 bat 用）
    print("MAKE_TABLE_L2 RESULT: %s  (mismatch=%d, x=%d)"
          % ("PASS" if (total_bad == 0 and total_x == 0) else "FAIL", total_bad, total_x))
    return 0 if (total_bad == 0 and total_x == 0) else 1


sys.exit(main())
