# -*- coding: utf-8 -*-
"""compare_python.py —— 把你自己的 Python 量化结果，和 RTL 仿真结果逐点对拍

背景：RTL 仿真结果已经和本工程的定点模型**逐点证明相等**
（tb_top_real：整帧 30720 个 plane unit 全对；3 个 tile 的 dwc/pwsum/qq/bnq/pool 0 处不一致）。
所以：**你的 Python 只要和这里的 golden 对上，就等于和 RTL 对上。**

用法（在工程根目录）：
    & 'D:\\...\\envs\\cyclegan\\python.exe' rtl\\conv2\\picture_and_para\\compare_python.py <你的结果文件>

支持格式（自动判定）：
    .npy                     形状 (120,160,8) / (8,120,160) / (160,120,8) / 拉平 153600
    .txt / .csv              空白或逗号分隔的数字，按上面任一顺序拉平
    .xlsx                    取第 1 个 sheet 里的全部数字，同上

数值口径：
    默认按 **Q4.4 整数** 比（和 golden 一样）。
    如果你的结果是**反量化后的实际值**（|x| ≤ 8 的小数），加 --deq，脚本会 ×16 取整再比。
    如果你的结果是 float32 但本来就是整数格点，不用加。

    --tol N    允许的绝对误差（默认 0 = 必须逐点相等）
    --out F    把差异写成 xlsx（默认 feature_maps_vs_python.xlsx）
"""
import os
import sys

import numpy as np

HERE = os.path.dirname(os.path.abspath(__file__))
GOLD_NPY = os.path.join(HERE, "golden_out_plane.npy")
REF_TXT = os.path.join(HERE, "golden_tiles.txt")

H, W, C = 120, 160, 8


# ----------------------------------------------------------------------
# 读入 + 形状归一
# ----------------------------------------------------------------------
def load_any(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".npy":
        return np.load(path)
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        vals = []
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, bool) or v is None:
                    continue
                if isinstance(v, (int, float)):
                    vals.append(v)
        return np.array(vals, dtype=np.float64)
    # txt / csv
    vals = []
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        for line in f:
            line = line.strip()
            if (not line) or line.startswith("#"):
                continue
            for tok in line.replace(",", " ").replace("\t", " ").split():
                try:
                    vals.append(float(tok))
                except ValueError:
                    pass
    return np.array(vals, dtype=np.float64)


def to_hwc(a):
    """归一成 (120,160,8)。返回 (数组, 说明) 或 (None, 原因)"""
    a = np.asarray(a)
    if a.size != H * W * C:
        return None, "元素个数 %d != %d (120*160*8)" % (a.size, H * W * C)
    if a.ndim == 3:
        if a.shape == (H, W, C):
            return a, "shape (120,160,8) 直接使用"
        if a.shape == (C, H, W):
            return np.transpose(a, (1, 2, 0)), "shape (8,120,160) → 转成 HWC"
        if a.shape == (W, H, C):
            return np.transpose(a, (1, 0, 2)), "shape (160,120,8) → 转成 HWC"
        if a.shape == (C, W, H):
            return np.transpose(a, (2, 1, 0)), "shape (8,160,120) → 转成 HWC"
    flat = a.reshape(-1)
    cand = [("按 (120,160,8) 拉平", flat.reshape(H, W, C)),
            ("按 (8,120,160) 拉平", flat.reshape(C, H, W).transpose(1, 2, 0)),
            ("按 (160,120,8) 拉平", flat.reshape(W, H, C).transpose(1, 0, 2))]
    if os.path.exists(GOLD_NPY):
        g = np.load(GOLD_NPY)
        best, bname, berr = None, None, None
        for name, arr in cand:
            e = np.abs(arr.astype(np.float64) - g).sum()
            if berr is None or e < berr:
                best, bname, berr = arr, name, e
        return best, "1-D %s（三种都试了，取最像的）" % bname
    return cand[0][1], "1-D 按 (120,160,8) 拉平"


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 1
    path = sys.argv[1]
    deq = "--deq" in sys.argv
    tol = 0
    outx = os.path.join(HERE, "feature_maps_vs_python.xlsx")
    for i, a in enumerate(sys.argv):
        if a == "--tol" and i + 1 < len(sys.argv):
            tol = int(sys.argv[i + 1])
        if a == "--out" and i + 1 < len(sys.argv):
            outx = sys.argv[i + 1]

    if not os.path.exists(GOLD_NPY):
        print("!! 没有 %s，先跑 gen_stim.py" % GOLD_NPY)
        return 1
    gold = np.load(GOLD_NPY).astype(np.float64)

    print("=" * 74)
    print("compare_python : 你的 Python 结果  vs  RTL 仿真 golden")
    print("=" * 74)
    raw = load_any(path)
    print("  读入 %s : dtype=%s size=%d" % (os.path.basename(path), raw.dtype, raw.size))

    got, how = to_hwc(raw)
    if got is None:
        print("  FAIL: %s" % how)
        return 1
    print("  形状归一: %s" % how)

    too_many_nonint = False
    if np.issubdtype(got.dtype, np.floating):
        frac = np.abs(got - np.round(got))
        too_many_nonint = float((frac > 0.02).mean()) > 0.5
    if deq or too_many_nonint:
        mx = float(np.abs(got).max())
        if mx <= 8.5 and (deq or too_many_nonint):
            print("  判定为**反量化实际值**（max|x|=%.3f ≤ 8.5）→ ×16 取整  [--deq]" % mx)
            got = np.round(got * 16.0)
        elif deq:
            print("  按 --deq ×16 取整（max|x|=%.3f）" % mx)
            got = np.round(got * 16.0)
    got = got.astype(np.int64)
    gold_i = gold.astype(np.int64)

    d = np.abs(got - gold_i)
    nbad = int((d > tol).sum())
    print("  总点数 = %d，不一致(>%d) = %d，最大绝对差 = %d" % (d.size, tol, nbad, int(d.max())))
    if nbad:
        idx = np.argwhere(d > tol)
        print("  前 20 个不一致 (y, x, oc): got vs exp")
        for k in idx[:20]:
            y, x, oc = int(k[0]), int(k[1]), int(k[2])
            print("     (%3d,%3d,oc%d)  got %5d  exp %5d   (实际值 %.3f vs %.3f)" %
                  (y, x, oc, got[y, x, oc], gold_i[y, x, oc], got[y, x, oc] / 16.0, gold_i[y, x, oc] / 16.0))
        try:
            from openpyxl import Workbook
            from openpyxl.styles import PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = "diff"
            ws.append(["y", "x", "oc", "你的值", "RTL/golden", "差", "你的实际值", "RTL 实际值"])
            for k in idx[:20000]:
                y, x, oc = int(k[0]), int(k[1]), int(k[2])
                ws.append([y, x, oc, int(got[y, x, oc]), int(gold_i[y, x, oc]),
                           int(d[y, x, oc]), got[y, x, oc] / 16.0, gold_i[y, x, oc] / 16.0])
            wb.save(outx)
            print("  差异明细写出: %s（最多 20000 条）" % outx)
        except Exception as e:                                   # noqa
            print("  （写 xlsx 失败：%s）" % e)

    print("\n---------------- 结论 ----------------")
    if nbad == 0:
        print("  MATCH：你的 Python 量化结果与 RTL 仿真**逐点一致**")
    else:
        per_oc = [int((d[:, :, c] > tol).sum()) for c in range(C)]
        print("  MISMATCH：%d/%d 点不一致（按 oc 分布 %s）" % (nbad, d.size, per_oc))
        print("  建议排查顺序：① 输入 Q4.4 口径 q=(p-124)>>>3 ② 权重 Q8 ③ BN 的 B_q 是否 ×4096")
        print("               ④ reflection pad 是 reflect-101 ⑤ dw/pw 两级 clamp(0,255) 的位置")
    print("--------------------------------------")
    return 0 if nbad == 0 else 1


sys.exit(main())
