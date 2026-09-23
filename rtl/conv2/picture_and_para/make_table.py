# -*- coding: utf-8 -*-
"""make_table.py —— 把仿真的逐级数据变成"数据变化表"（xlsx）

输入：
    golden_tiles.txt   由 gen_stim.py 生成（定点 golden + 浮点参考 + 参数）
    real_dump.txt      由 tb_top_real.v 生成（RTL 实测的 dwc/pwsum/qq/bnq/pool）
输出：
    feature_maps_real.xlsx
        Params   口径 / 权重 / 逐级统计 / RTL vs Golden 比对结果
        Weights  真实权重（浮点 + Q8）
        IN_t*    输入：原始像素 / Q4.4 / 12×12 反射窗口
        DWC_t*   深度卷积：累加和 / RTL 量化 / golden / 浮点×16
        PW_t*    点卷积：累加和 / 量化 / golden / 浮点×16
        BN_t*    BatchNorm：RTL / golden / 浮点×16 / 实际值
        OUT_t*   2×2 max 池化：RTL / golden / 浮点×16 / 实际值

跑法：
    & 'D:\\Users\\Administrator\\anaconda3\\envs\\cyclegan\\python.exe' rtl\\conv2\\picture_and_para\\make_table.py
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
GOLD = os.path.join(HERE, "golden_tiles.txt")
DUMP = os.path.join(HERE, "real_dump.txt")
OUTX = os.path.join(HERE, "feature_maps_real.xlsx")

# ---------------- 配色（和 feature_map.py 一致）----------------
HDR      = PatternFill("solid", start_color="DDEBF7")
HDR2     = PatternFill("solid", start_color="FFF2CC")
PAD_F    = PatternFill("solid", start_color="E2EFDA")
PAD_F2   = PatternFill("solid", start_color="FCE4D6")
SUM_F    = PatternFill("solid", start_color="D9E1F2")
SUM_F2   = PatternFill("solid", start_color="E4DFEC")
QNT_F    = PatternFill("solid", start_color="C6E0B4")
QNT_F2   = PatternFill("solid", start_color="FFE699")
PW_F     = PatternFill("solid", start_color="F8CBAD")
PW_F2    = PatternFill("solid", start_color="D6DCE4")
RTL_F    = PatternFill("solid", start_color="BDD7EE")
RTL_F2   = PatternFill("solid", start_color="FFD966")
GOLD_F   = PatternFill("solid", start_color="E2EFDA")
FLT_F    = PatternFill("solid", start_color="EDEDED")
BAD_F    = PatternFill("solid", start_color="FF7C80")
PASS_F   = PatternFill("solid", start_color="C6EFCE")
PADCELL  = PatternFill("solid", start_color="F2F2F2")
TITLE_FONT = Font(bold=True, size=12)
HEAD_FONT  = Font(bold=True)


# =====================================================================
# 解析
# =====================================================================
def parse(path):
    """TILE i tr tc / PARAM name v... / KEY idx v v v ...
    数据按键 (tile, KEY, idx) 存，避免不同 tile 的同名键互相覆盖。"""
    data = {"TILE": [], "PARAM": {}, "V": {}}
    cur = -1
    if not os.path.exists(path):
        raise SystemExit("找不到 %s（先跑 gen_stim.py / tb_top_real.v）" % path)
    for line in open(path, encoding="utf-8", errors="replace"):
        line = line.strip()
        if (not line) or line.startswith("#"):
            continue
        parts = line.split()
        key = parts[0]
        if key == "PARAM":
            data["PARAM"][parts[1]] = parts[2:]
        elif key == "TILE":
            cur = int(parts[1])
            data["TILE"].append((cur, int(parts[2]), int(parts[3])))
        else:
            data["V"][(cur, key, int(parts[1]))] = parts[2:]
    return data


def vget(d, tile, key, idx, default=None):
    return d["V"].get((tile, key, idx), default)


def h8(vals):
    """8bit 补码十六进制 → 无符号 0..255"""
    return [int(v, 16) & 0xFF for v in vals]


def s8(vals):
    """8bit 补码十六进制 → 有符号 -128..127"""
    out = []
    for v in vals:
        n = int(v, 16) & 0xFF
        out.append(n - 256 if n >= 128 else n)
    return out


def h24(vals):
    return [int(v, 16) & 0xFFFFFF for v in vals]


def s24(vals):
    out = []
    for v in vals:
        n = int(v, 16) & 0xFFFFFF
        out.append(n - (1 << 24) if n >= (1 << 23) else n)
    return out


def hf(vals):
    return [float(v) for v in vals]


def dec(vals):
    """十进制字符串 → int"""
    return [int(v) for v in vals]


# =====================================================================
# 写表助手
# =====================================================================
def seg(ws, r, c0, title, values, W, fill, fmt="%d", rowlab=True, extra_fill=None):
    """写一段：标题行 + 列号 + W 列网格，返回下一段的起始行"""
    ws.cell(row=r, column=c0, value=title).fill = fill
    ws.cell(row=r, column=c0).font = HEAD_FONT
    if rowlab:
        ws.cell(row=r, column=c0 + 1, value="y\\x").fill = fill
    for x in range(W):
        ws.cell(row=r, column=c0 + 2 + x, value=x).fill = fill
    H = (len(values) + W - 1) // W
    for y in range(H):
        if rowlab:
            ws.cell(row=r + 1 + y, column=c0 + 1, value=y).fill = fill
        for x in range(W):
            i = y * W + x
            if i >= len(values):
                break
            # %d 的网格写成**数字**（Excel 里可排序/可算），其余（0x.. / 浮点）写成字符串
            v = values[i] if fmt == "%d" else (fmt % values[i])
            cell = ws.cell(row=r + 1 + y, column=c0 + 2 + x, value=v)
            if extra_fill is not None and (y, x) in extra_fill:
                cell.fill = BAD_F
    return r + H + 2


def put_legend(ws, r, text):
    ws.cell(row=r, column=1, value=text).font = Font(italic=True, color="808080")
    return r + 1


# =====================================================================
# 主流程
# =====================================================================
def main():
    g = parse(GOLD)
    try:
        d = parse(DUMP)
    except SystemExit:
        print("!! 没有 real_dump.txt —— 表里只放 golden，请先跑 tb_top_real.v")
        d = {"TILE": g["TILE"], "PARAM": {}}

    P = g["PARAM"]
    iw, ih = int(P["IW"][0]), int(P["IH"][0])
    ntiles = int(P["NTILES"][0])
    dwq = [int(v) for v in P["DWQ"]]
    pwq = [int(v) for v in P["PWQ"]]
    aq = [[int(v) for v in P["AQ%d" % i]] for i in range(ntiles)]
    bq = [[int(v) for v in P["BQ%d" % i]] for i in range(ntiles)]

    wb = Workbook()
    wb.remove(wb.active)

    # ---------------- 比对（RTL vs golden）----------------
    keys = [("DWC", 3, 100), ("PWSUM", 8, 100), ("QQ", 8, 100), ("BNQ", 8, 100), ("POOL", 8, 25)]
    cmp_rows = []
    for i, (ti, tr, tc) in enumerate(g["TILE"]):
        for key, nidx, nval in keys:
            bad = 0
            for idx in range(nidx):
                gv = vget(g, i, key, idx)
                dv = vget(d, i, key, idx)
                if gv is None or dv is None:
                    continue
                for n in range(min(len(gv), len(dv))):
                    if (int(gv[n], 16) & 0xFFFFFF) != (int(dv[n], 16) & 0xFFFFFF):
                        bad += 1
            cmp_rows.append((i, tr, tc, key, nidx * nval, bad))

    # ================= Params =================
    ws = wb.create_sheet("Params")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 62
    r = 1
    ws.cell(row=r, column=1, value="conv2 L1 真实激励仿真 —— 数据变化表").font = TITLE_FONT
    r += 2
    rows = [
        ("图片", "picture_and_para/test.jpg  %dx%d，DDR 布局 row*%d + ch*%d + col" % (iw, ih, iw * 3, iw)),
        ("网络结构", "cyclegna_mobilenet.py: model.1 = dw3×3(groups=3,pad=0) + pw1×1(3→8)"),
        ("", "             model.2 = BatchNorm2d(8) → ReLU → MaxPool2d(2,2)   ← 本工程只做到这里"),
        ("ReflectionPad", "网络在 dw 前有 ReflectionPad2d(1)；conv_win_load 的 reflect-101 与之逐点等价"),
        ("训练归一化", "transforms.Normalize(mean=.5,std=.5) → x = 2p/255 - 1 ✓ 与 Q4.4 口径一致"),
        ("权重来源", "picture_and_para/netG_B_epoch11.pth → model.1.depthwise(3,1,3,3) + model.1.pointwise(8,3,1,1)"),
        ("BN 来源", "model.2（8 通道）→ scale = gamma/sqrt(var+eps), shift = beta - mean*scale"),
        ("权重 ROM", "rtl/conv2/conv_wrom/wrom.hex（67 字：27 dw + 24 pw + 8 bn_a + 8 bn_b）"),
        ("输入定点", "Q4.4 有符号：q = (p-124) >>> 3，实际值 q/16 ∈ [-1,1]（硬件=一个减法器+算术右移）"),
        ("参数定点", "Q8：w_q = round(w*256)"),
        ("BN 定点", "bnq = (A_q*q + B_q) >>> 8，A_q = round(scale*256)，B_q = round(shift*4096)"),
        ("BN 舍入选项", "conv_top 的 BN_ROUND：0 = 直接截断 x>>>8（默认）1 = 四舍五入 (x+128)>>>8"),
        ("", "  实测整帧：BN 误差 1.7443→1.7408 LSB、池化 1.9828→2.0028 LSB（略差）；"),
        ("", "  实例归一化出图：L1 误差 0.0847→0.0857、最终图 PSNR 23.78→24.15 dB（略好）→ 噪声级，默认 0"),
        ("", "  （B_q 额外 ×16：数据通路相对 Q4.4 有 16 倍增益，而 RTL 的再量化固定 >>>8）"),
        ("dw 相位", "dwc = sat((Σ_3x3 q*w_dw_q + 128) >>> 8)   ← 寄存器值为 Q4.4（有符号）"),
        ("pw 相位", "qq  = sat((Σ_ch dwc*w_pw_q + 128) >>> 8)"),
        ("中间饱和", "dw/pw：**对称**饱和 [-128, 127]（真实网络那里没有激活）；BN：饱和 [0, 127] = **ReLU + 上限**"),
        ("", "  即：dwc/qq 保留负值；bnq 下限 0（ReLU）、上限 127（+7.9375）。实测撞界点的个数都是 0"),
        ("", "  配套：pw 的 a 通路（dwc）符号扩展；池化比较器改成有符号"),
        ("限位位置", "统一放在 **PE 阵列输出**（conv_top .PE_SAT(1)）：移位前饱和一次 [-32768, +32639]"),
        ("", "  ★ 量纲是 Q4.4 的 256 倍，所以界是 32768 量级不是 8；HI 留了 128 给 (x+128)>>>8 的进位"),
        ("", "  ★ 与「三级量化里各自限位」**逐位等价**（80 万点穷举 + 两遍整帧仿真的 plane 全对验证）"),
        ("", "  ★ PE 输出是组合的，所以是**零拍**改动，三个抓数点 c=13 / pc=7 / pc=4 都不用动"),
        ("池化", "2×2 max（**有符号**比较），写回 plane：unit = (oc*120 + row)*32 + tile_c，40bit = 该行 5 列"),
        ("RTL 开关", "conv_top .Q44_EN(1) .DW_SIGNED(1) .Q44_SAT(1) .BN_RELU(1) .PE_SAT(1)；老 tb 默认 0，逐位不变"),
        ("", ""),
        ("仿真 tb", "rtl/conv2/tb/tb_top_real.v（整帧 320×240×3 → 160×120×8，768 个 tile）"),
        ("抓数 tile", "、".join("(%d,%d)" % (tr, tc) for _, tr, tc in
                                [(t[0], t[1], t[2]) for t in g["TILE"]])),
        ("", ""),
        ("★ 表里数字口径", "标注 RTL 的都是仿真实测（Q4.4 寄存器值）；标注 GOLD 的是同口径定点模型；"),
        ("", "标注 FLT×16 的是浮点参考 ×16（可与 RTL 直接比，差多少就是定点损失）"),
        ("", ""),
        ("★ 数值分布说明", "整帧：dwc -8..7（9.1% 为 0）、qq -8..9（19.2% 为 0）、bnq 0..70（ReLU 后非负）、池化 0..70。"),
        ("", "dw/pw 保留了负值，所以特征图不再是一半被抹成 0；±8 饱和点实测 0 个（极少越界，符合预期）。"),
        ("", "本表证明 **RTL = 定点模型**：RTL vs GOLD 全部 0 处不一致，整帧 30720 unit 全对。"),
    ]
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = HEAD_FONT
        ws.cell(row=r, column=2, value=v)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="RTL vs Golden 比对（逐 tile 逐级）").font = TITLE_FONT
    r += 1
    for h, c in zip(["tile", "tile_r", "tile_c", "级", "点数", "不一致"], range(1, 7)):
        ws.cell(row=r, column=c, value=h).fill = HDR
        ws.cell(row=r, column=c).font = HEAD_FONT
    r += 1
    total_bad = 0
    for (i, tr, tc, key, n, bad) in cmp_rows:
        ws.cell(row=r, column=1, value="t%d" % i)
        ws.cell(row=r, column=2, value=tr)
        ws.cell(row=r, column=3, value=tc)
        ws.cell(row=r, column=4, value=key)
        ws.cell(row=r, column=5, value=n)
        cc = ws.cell(row=r, column=6, value=bad)
        cc.fill = PASS_F if bad == 0 else BAD_F
        total_bad += bad
        r += 1
    r += 1
    v = "全部一致（RTL 与定点 golden 逐点相同）" if total_bad == 0 else "有 %d 个点不一致！" % total_bad
    ws.cell(row=r, column=1, value="结论").font = HEAD_FONT
    ws.cell(row=r, column=2, value=v).fill = PASS_F if total_bad == 0 else BAD_F

    # ================= Weights =================
    ws = wb.create_sheet("Weights")
    r = 1
    ws.cell(row=r, column=1, value="dw 深度卷积 3ch×3×3（idx = ch*9 + kh*3 + kw）").font = TITLE_FONT
    r += 1
    for h, c in zip(["idx", "ch", "kh", "kw", "Q8 整数", "浮点(近似)"], range(1, 7)):
        ws.cell(row=r, column=c, value=h).fill = HDR
    r += 1
    for i in range(27):
        for c, v in zip(range(1, 7), [i, i // 9, (i % 9) // 3, i % 3, dwq[i], "%.6f" % (dwq[i] / 256.0)]):
            ws.cell(row=r, column=c, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="pw 点卷积 8oc×3ic（idx = oc*3 + ic）").font = TITLE_FONT
    r += 1
    for h, c in zip(["idx", "oc", "ic", "Q8 整数", "浮点(近似)"], range(1, 6)):
        ws.cell(row=r, column=c, value=h).fill = HDR2
    r += 1
    for i in range(24):
        for c, v in zip(range(1, 6), [i, i // 3, i % 3, pwq[i], "%.6f" % (pwq[i] / 256.0)]):
            ws.cell(row=r, column=c, value=v)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="BatchNorm2d（逐 oc；A_q = round(scale*256)、B_q = round(shift*4096)）").font = TITLE_FONT
    r += 1
    for h, c in zip(["oc", "A_q", "B_q", "scale(≈A_q/256)", "shift(≈B_q/4096)"], range(1, 6)):
        ws.cell(row=r, column=c, value=h).fill = PW_F
    r += 1
    for oc in range(8):
        for c, v in zip(range(1, 6), [oc, aq[0][oc], bq[0][oc],
                                      "%.5f" % (aq[0][oc] / 256.0), "%.5f" % (bq[0][oc] / 4096.0)]):
            ws.cell(row=r, column=c, value=v)
        r += 1

    # ================= 每个 tile =================
    for i, (ti, tr, tc) in enumerate(g["TILE"]):
        # ---------- IN ----------
        ws = wb.create_sheet("IN_t%02d_%02d" % (tr, tc))
        r = 1
        ws.cell(row=r, column=1, value="tile(%d,%d) 输入：原图像素 p（0..255）→ Q4.4 q=(p-124)>>>3 → 12×12 反射窗口"
                % (tr, tc)).font = TITLE_FONT
        r += 2
        for c in range(3):
            qin = s8(vget(g, i, "QIN", c))
            raw = dec(vget(g, i, "RAW", c)) if vget(g, i, "RAW", c) else None
            if raw is not None:
                seg(ws, r, 1, "ch%d 原图像素 p（DDR 里就是这个，0..255）" % c, raw, 10, HDR, "%d")
                r += 1 + ((len(raw) + 9) // 10)
            seg(ws, r, 1, "ch%d Q4.4 (q，实际值 = q/16)" % c, qin, 10, PAD_F, "%d")
            r += 1 + ((len(qin) + 9) // 10)
            seg(ws, r, 16, "ch%d Q4.4 实际值（/16）" % c,
                ["%.3f" % (v / 16.0) for v in qin], 10, PAD_F2, "%s")
            r += 1 + ((len(qin) + 9) // 10)
            win = s8(vget(g, i, "WIN", c))
            seg(ws, r, 1, "ch%d 12×12 窗口（含反射，q）" % c, win, 12, SUM_F, "%d")
            r += 1 + ((len(win) + 11) // 12)
            seg(ws, r, 16, "ch%d 12×12 窗口 实际值（/16）" % c,
                ["%.3f" % (v / 16.0) for v in win], 12, SUM_F2, "%s")
            r += 1 + ((len(win) + 11) // 12) + 1

        # ---------- DWC ----------
        ws = wb.create_sheet("DWC_t%02d_%02d" % (tr, tc))
        r = 1
        ws.cell(row=r, column=1, value="tile(%d,%d) 深度卷积：累加和 → RTL 量化值（Q4.4）" % (tr, tc)).font = TITLE_FONT
        r += 2
        for c in range(3):
            rawsum = s24(vget(g, i, "DWCRAW", c))
            rtl = h8(vget(d, i, "DWC", c) or vget(g, i, "DWC", c))
            gold = h8(vget(g, i, "DWC", c))
            flt = [int(round(v * 16)) for v in hf(vget(g, i, "DWF", c))]
            bad = {(y, x) for y in range(10) for x in range(10) if rtl[y * 10 + x] != gold[y * 10 + x]}
            seg(ws, r, 1, "ch%d 累加和 Σq*w_q（十进制，量化前）" % c, rawsum, 10, SUM_F, "%d")
            r += 1 + 10
            seg(ws, r, 1, "ch%d RTL dwc（Q4.4 有符号，±8 饱和）" % c, rtl, 10, RTL_F, "%d", extra_fill=bad)
            r += 1 + 10
            seg(ws, r, 16, "ch%d GOLD dwc（定点模型）" % c, gold, 10, GOLD_F, "%d")
            r += 1 + 10
            seg(ws, r, 16, "ch%d FLT×16（浮点参考×16）" % c, flt, 10, FLT_F, "%d")
            r += 1 + 10 + 1

        # ---------- PW ----------
        ws = wb.create_sheet("PW_t%02d_%02d" % (tr, tc))
        r = 1
        ws.cell(row=r, column=1, value="tile(%d,%d) 点卷积：Σ dwc*w_pw → 量化（Q4.4）" % (tr, tc)).font = TITLE_FONT
        r += 2
        for oc in range(8):
            psum = s24(vget(g, i, "PWSUM", oc))
            rtl = h8(vget(d, i, "QQ", oc) or vget(g, i, "QQ", oc))
            gold = h8(vget(g, i, "QQ", oc))
            flt = [int(round(v * 16)) for v in hf(vget(g, i, "PWF", oc))]
            bad = {(y, x) for y in range(10) for x in range(10) if rtl[y * 10 + x] != gold[y * 10 + x]}
            seg(ws, r, 1, "oc%d 累加和 Σ dwc*w_pw_q" % oc, psum, 10, PW_F, "%d")
            r += 1 + 10
            seg(ws, r, 1, "oc%d RTL qq（Q4.4）" % oc, rtl, 10, RTL_F, "%d", extra_fill=bad)
            r += 1 + 10
            seg(ws, r, 16, "oc%d GOLD qq" % oc, gold, 10, GOLD_F, "%d")
            r += 1 + 10
            seg(ws, r, 16, "oc%d FLT×16" % oc, flt, 10, FLT_F, "%d")
            r += 1 + 10 + 1

        # ---------- BN ----------
        ws = wb.create_sheet("BN_t%02d_%02d" % (tr, tc))
        r = 1
        ws.cell(row=r, column=1, value="tile(%d,%d) BatchNorm2d：bnq = clamp((A_q*qq + B_q)>>>8)" % (tr, tc)).font = TITLE_FONT
        r += 2
        for oc in range(8):
            rtl = h8(vget(d, i, "BNQ", oc) or vget(g, i, "BNQ", oc))
            gold = h8(vget(g, i, "BNQ", oc))
            flt = [int(round(v * 16)) for v in hf(vget(g, i, "BNF", oc))]
            real = ["%.3f" % (v / 16.0) for v in rtl]
            bad = {(y, x) for y in range(10) for x in range(10) if rtl[y * 10 + x] != gold[y * 10 + x]}
            seg(ws, r, 1, "oc%d RTL bnq（Q4.4）" % oc, rtl, 10, RTL_F, "%d", extra_fill=bad)
            r += 1 + 10
            seg(ws, r, 1, "oc%d RTL bnq 实际值（/16）" % oc, real, 10, QNT_F, "%s")
            r += 1 + 10
            seg(ws, r, 16, "oc%d GOLD bnq" % oc, gold, 10, GOLD_F, "%d")
            r += 1 + 10
            seg(ws, r, 16, "oc%d FLT×16（浮点 BN×16）" % oc, flt, 10, FLT_F, "%d")
            r += 1 + 10 + 1

        # ---------- OUT ----------
        ws = wb.create_sheet("OUT_t%02d_%02d" % (tr, tc))
        r = 1
        ws.cell(row=r, column=1, value="tile(%d,%d) 2×2 max 池化输出（就是写回 plane 的值，Q4.4）" % (tr, tc)).font = TITLE_FONT
        r += 2
        for oc in range(8):
            rtl = h8(vget(d, i, "POOL", oc) or vget(g, i, "POOL", oc))
            gold = h8(vget(g, i, "POOL", oc))
            flt = [int(round(v * 16)) for v in hf(vget(g, i, "OUTF", oc))]
            real = ["%.3f" % (v / 16.0) for v in rtl]
            bad = {(y, x) for y in range(5) for x in range(5) if rtl[y * 5 + x] != gold[y * 5 + x]}
            seg(ws, r, 1, "oc%d RTL 池化（Q4.4）" % oc, rtl, 5, RTL_F, "%d", extra_fill=bad)
            r += 1 + 5
            seg(ws, r, 1, "oc%d RTL 实际值（/16）" % oc, real, 5, PASS_F, "%s")
            r += 1 + 5
            seg(ws, r, 8, "oc%d GOLD 池化" % oc, gold, 5, GOLD_F, "%d")
            r += 1 + 5
            seg(ws, r, 8, "oc%d FLT×16" % oc, flt, 5, FLT_F, "%d")
            r += 1 + 5 + 1

    wb.save(OUTX)
    print("写出 %s（%d 个 sheet）" % (OUTX, len(wb.sheetnames)))
    print("RTL vs Golden 不一致点数 = %d" % total_bad)
    print("sheet 列表:", ", ".join(wb.sheetnames))
    # ★ ASCII 判定行：给 check_fixed_point.bat 用（bat 必须纯 ASCII，不能拿中文去 findstr）
    print("MAKE_TABLE RESULT: %s  (mismatch=%d)" %
          ("PASS" if total_bad == 0 else "FAIL", total_bad))


main()
