# -*- coding: utf-8 -*-
"""对拍：gen_stim 从 .pth 解析出的权重/BN，与你 netG_B_epoch11_weights.xlsx 逐个数字比。

注意两个展平顺序不同（这是刚才"看起来不一致"的唯一原因）：
  · 你的 xlsx model.1.depthwise：按 **kh 行** 排，每行给出 ch0/ch1/ch2 的 3 个值
      r2 = [ch0r0 x3][空][ch1r0 x3][空][ch2r0 x3]
  · 工程的 w_dw[] 顺序：按 **通道** 排 —— w_dw[ch*9 + kh*3 + kw]
"""
import os
import numpy as np
from openpyxl import load_workbook

import stim_model as M

W = M.load_weights()
wb = load_workbook(os.path.join(M.HERE, "netG_B_epoch11_weights.xlsx"), data_only=True, read_only=True)


def fl(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


rows = [list(r) for r in wb["model.1"].iter_rows(values_only=True)]

# ---- dw：r2..r4，每行 3 个块（ch0/ch1/ch2），块内 3 个 = (kw 0..2) ----
sheet_dw = np.zeros((3, 3, 3))
for kh in range(3):
    row = rows[2 + kh]
    for ch in range(3):
        blk = [fl(row[ch * 4 + kw]) for kw in range(3)]
        for kw in range(3):
            sheet_dw[ch, kh, kw] = blk[kw]

# ---- pw：r8, r10, ..., r22，每个 oc 一行 3 个 ----
sheet_pw = np.zeros((8, 3))
for oc in range(8):
    row = rows[8 + oc * 2]
    for ic in range(3):
        sheet_pw[oc, ic] = fl(row[ic])

# ---- BN：model.2 的 a/b 列 ----
rows2 = [list(r) for r in wb["model.2"].iter_rows(values_only=True)]
sheet_a = np.zeros(8)
sheet_b = np.zeros(8)
for oc in range(8):
    r = rows2[2 + oc]
    sheet_a[oc] = fl(r[6])
    sheet_b[oc] = fl(r[7])

print("=" * 78)
print("权重/BN 对拍（你的 xlsx  vs  我从 .pth 解析）")
print("=" * 78)
d1 = np.abs(sheet_dw - W["dw"])
d2 = np.abs(sheet_pw - W["pw"])
d3 = np.abs(sheet_a - W["scale"])
d4 = np.abs(sheet_b - W["shift"])
print("  dw 3x3x3 (27) : 最大绝对差 = %.3e" % d1.max())
print("  pw 8x3   (24) : 最大绝对差 = %.3e" % d2.max())
print("  BN a = scale  : 最大绝对差 = %.3e" % d3.max())
print("  BN b = shift  : 最大绝对差 = %.3e" % d4.max())
allok = max(d1.max(), d2.max(), d3.max(), d4.max()) < 5e-6
print("  结论：%s" % ("逐点一致（差异只是 xlsx 里显示的位数）" if allok else "有差异！"))

print("\n  逐通道 dw 第一行对照：")
for ch in range(3):
    print("     ch%d  表=%s  我=%s" %
          (ch, np.round(sheet_dw[ch, 0], 6).tolist(), np.round(W["dw"][ch, 0], 6).tolist()))
print("  pw 前 3 个 oc：")
for oc in range(3):
    print("     oc%d  表=%s  我=%s" % (oc, np.round(sheet_pw[oc], 6).tolist(), np.round(W["pw"][oc], 6).tolist()))
print("  BN  a/b：")
for oc in range(3):
    print("     oc%d  表 a=%.5f b=%.5f   我 scale=%.5f shift=%.5f" %
          (oc, sheet_a[oc], sheet_b[oc], W["scale"][oc], W["shift"][oc]))

print("\n  量化后的定点参数（ROM 里就是这些）")
print("     Q8 dw =", list(W["dw_q"].reshape(-1)))
print("     Q8 pw =", list(W["pw_q"].reshape(-1)))
print("     A_q   =", list(W["a_q"]))
print("     B_q   =", list(W["b_q"]))
